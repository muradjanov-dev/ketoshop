"""Finance Excel export keeps booking values and delivery-period profit distinct."""
import asyncio
import json
import os
import unittest
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

os.environ.setdefault("DATABASE_URL", "postgresql://test/test")
os.environ.setdefault("BOT_TOKEN", "1:test")

from openpyxl import load_workbook

import reports
import admin_web


class FinanceExcelReportTest(unittest.TestCase):
    def test_order_total_appears_once_and_summary_uses_delivery_period_stats(self):
        orders = [
            {
                "id": 101, "created_at": None, "status": "pending", "total": 250,
                "items_data": [
                    {"name": "A", "product_id": 10, "quantity": 1, "price": 100},
                    {"name": "B", "product_id": 10, "quantity": 1, "price": 150},
                ],
            },
            {
                "id": 102, "created_at": None, "status": "cancelled", "total": 90,
                "items_data": [{"name": "C", "product_id": 10, "quantity": 1, "price": 90}],
            },
        ]
        stats = {
            "orders_sold": 1, "booked_value": 250,
            "orders_delivered": 1, "delivered_revenue": 225,
            "product_cost": 80, "expenses": 25, "profit": 120,
        }

        with patch.object(reports, "get_orders_for_export", AsyncMock(return_value=orders)), \
             patch.object(reports, "get_admin_stats", AsyncMock(return_value=stats)), \
             patch.object(reports, "get_all_cost_prices", AsyncMock(return_value={10: 80})), \
             patch.object(reports, "get_set_costs", AsyncMock(return_value={})), \
             patch.object(reports, "_fmt_dt", return_value=""):
            buffer, _ = asyncio.run(reports.generate_orders_excel("30d"))

        workbook = load_workbook(BytesIO(buffer.getvalue()), data_only=True)
        detail = workbook["Buyurtmalar"]
        self.assertEqual(detail.cell(row=2, column=11).value, 250)
        self.assertIsNone(detail.cell(row=3, column=11).value)
        self.assertIsNone(detail.cell(row=4, column=11).value)  # cancelled order

        summary = workbook["Xulosa"]
        amounts = {summary.cell(row=r, column=1).value: summary.cell(row=r, column=2).value
                   for r in range(1, summary.max_row + 1)}
        self.assertEqual(amounts["Buyurtma qiymati (tushumga yozilgan)"], 250)
        self.assertEqual(amounts["Yetkazilgan tushum"], 225)
        self.assertEqual(amounts["Davr xarajatlari"], 25)
        self.assertEqual(amounts["Sof foyda (yetkazilgan tushum − tannarx − xarajat)"], 120)


class FinanceDashboardPayloadTest(unittest.TestCase):
    def test_dashboard_keeps_booked_and_delivered_values_separate(self):
        stats = {
            "orders_sold": 4, "booked_value": 900,
            "orders_delivered": 2, "delivered_revenue": 420,
            "revenue": 420, "product_cost": 190, "expenses": 30, "profit": 200,
            "b2b_revenue": 100, "b2b_orders": 1, "orders_total": 5,
            "orders_pending": 1, "orders_confirmed": 0, "orders_cancelled": 0,
            "keto_discount": 0,
        }
        request = SimpleNamespace(
            cookies={admin_web.SESSION_COOKIE: admin_web._make_session()},
            query={"period": "30d"},
        )
        with patch.object(admin_web, "ADMIN_WEB_PASSWORD", "test"), \
             patch.object(admin_web.database, "get_admin_stats", AsyncMock(return_value=stats)), \
             patch.object(admin_web.database, "get_monthly_breakdown", AsyncMock(return_value=[])), \
             patch.object(admin_web.database, "get_top_products", AsyncMock(return_value=[])), \
             patch.object(admin_web.database, "get_keto_program_stats", AsyncMock(return_value={})), \
             patch.object(admin_web.database, "get_abc_analysis", AsyncMock(return_value=[])), \
             patch.object(admin_web.database, "get_b2b_orders", AsyncMock(return_value=[])), \
             patch.object(admin_web.database, "get_orders_by_region", AsyncMock(return_value=[])):
            response = asyncio.run(admin_web.api_dashboard(request))

        payload = json.loads(response.text)
        self.assertEqual(payload["stats"]["booked_value"], 900)
        self.assertEqual(payload["stats"]["orders_sold"], 4)
        self.assertEqual(payload["stats"]["delivered_revenue"], 420)
        self.assertEqual(payload["stats"]["orders_delivered"], 2)

        dashboard = (Path(__file__).parents[1] / "webapp" / "admin.html").read_text()
        self.assertIn("fmt(st.booked_value)", dashboard)
        self.assertIn("fmt(st.delivered_revenue)", dashboard)
        self.assertIn("st.orders_sold", dashboard)
        self.assertIn("st.orders_delivered", dashboard)
        self.assertIn("r.delivered_revenue", dashboard)
        self.assertIn("r.delivered_orders", dashboard)


if __name__ == "__main__":
    unittest.main()
