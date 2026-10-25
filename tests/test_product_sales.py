"""Dashboard «📦 Mahsulot savdosi», ▲/▼ against the previous window, and the
daily chart (owner, 2026-10-24). A fake connection — no database."""
import asyncio
import contextlib
import json
import os
import unittest
from datetime import date, datetime, timedelta
from io import BytesIO
from unittest.mock import AsyncMock, patch

os.environ.setdefault("DATABASE_URL", "postgresql://test/test")
os.environ.setdefault("BOT_TOKEN", "1:test")

import database
import reports


def run(coro):
    return asyncio.run(coro)


class _FakePool:
    def __init__(self, conn):
        self._conn = conn

    def acquire(self):
        @contextlib.asynccontextmanager
        async def _cm():
            yield self._conn
        return _cm()


ORDERS = [
    {"id": 1, "source": "bot", "created_at": datetime(2026, 10, 12, 8, 0), "items": json.dumps([
        {"product_id": 10, "name": "Eritritol 500 gr", "quantity": 2, "price": 60_000},
        {"product_id": 10, "name": "Eritritol 500 gr", "quantity": 1, "price": 0, "is_gift": True},
        {"product_id": 20, "name": "Kokos uni", "quantity": 1, "price": 45_000},
    ])},
    {"id": 2, "source": "b2b", "created_at": datetime(2026, 10, 20, 8, 0), "items": [
        {"product_id": 10, "name": "Eritritol 500 gr", "quantity": 10, "price": 50_000},
    ]},
    {"id": 3, "source": "bot", "created_at": datetime(2026, 10, 21, 8, 0), "items": json.dumps([
        {"is_set": True, "set_id": 7, "name": "Keto start", "quantity": 1, "price": 150_000},
        {"product_id": 30, "name": "Stevia", "quantity": 1, "price": 30_000},
        {"product_id": 99999, "name": "Old product", "quantity": 1, "price": 10_000},
    ])},
    {"id": 4, "source": "bot", "created_at": datetime(2026, 10, 22, 8, 0), "items": "not json"},
]
PRODUCTS = [
    {"id": 10, "name": "Eritritol 500 gr", "unit": "dona", "quantity": 26, "cost_price": 30_000,
     "category": "shakar", "is_active": 1},
    {"id": 20, "name": "Kokos uni", "unit": "dona", "quantity": 0, "cost_price": None,
     "category": "un", "is_active": 1},
    {"id": 30, "name": "Stevia", "unit": "dona", "quantity": 5, "cost_price": 12_000,
     "category": "shakar", "is_active": 1},
]


class _SalesConn:
    def __init__(self):
        self.queries = []

    async def fetch(self, sql, *args):
        self.queries.append((sql, args))
        if "FROM orders" in sql:
            return ORDERS
        if "FROM products" in sql:
            return PRODUCTS
        if "FROM product_sets" in sql:
            return [{"id": 7, "name": "Keto start to'plami"}]
        return []


class ProductSalesTest(unittest.TestCase):
    def test_postage_uses_saved_order_charge_without_repricing_old_sales(self):
        base = {"source": "bot", "created_at": datetime(2026, 10, 24, 12, 0),
                "status": "delivered", "keto_redeemed": 0}
        goods = [{"product_id": 10, "name": "Eritritol 500 gr", "quantity": 1, "price": 60_000},
                 {"product_id": 20, "name": "Kokos uni", "quantity": 1, "price": 40_000}]
        orders = [
            dict(base, id=1, delivery_method="uzpost", total=113_000, keto_redeemed=2_000,
                 items=json.dumps(goods)),  # 100 000 goods + 15 000 post - 2 000 Keto
            dict(base, id=2, delivery_method="uzpost", total=115_000, keto_redeemed=3_000,
                 created_at=datetime(2026, 10, 24, 14, 0), items=goods),  # saved 18 000 post
            dict(base, id=3, delivery_method="bts", total=100_000, items=goods),
            dict(base, id=4, delivery_method="self", total=125_000, items=goods),
            dict(base, id=5, delivery_method="uzpost", total=18_000, items="not json"),
            dict(base, id=6, delivery_method="uzpost", total=18_000, items=[]),
            dict(base, id=7, delivery_method="uzpost", total=90_000, items=goods),
        ]
        with patch(__name__ + ".ORDERS", orders):
            data, conn = self._run({"start": "2026-10-10", "end": "2026-10-24"})
        self.assertEqual(data.get("postage_orders", []), [
            {"id": 1, "created_at": datetime(2026, 10, 24, 12, 0), "method": "uzpost", "fee": 15_000},
            {"id": 2, "created_at": datetime(2026, 10, 24, 14, 0), "method": "uzpost", "fee": 18_000},
        ])
        self.assertEqual(data["totals"]["revenue"], 500_000)
        self.assertIn(database.SALE_SQL, conn.queries[0][0])
        self.assertIn("created_at >=", conn.queries[0][0])
        self.assertIn("created_at <=", conn.queries[0][0])

    def _run(self, period):
        conn = _SalesConn()
        with patch.object(database, "pool", _FakePool(conn)), \
                patch.object(database, "get_set_costs", AsyncMock(return_value={7: 90_000})):
            return run(database.get_product_sales(period)), conn

    def test_rows_money_and_flags(self):
        data, conn = self._run({"start": "2026-10-10", "end": "2026-10-24"})
        sql = conn.queries[0][0]
        self.assertIn(database.SALE_SQL, sql)
        self.assertIn("created_at >=", sql)
        rows = {(r["kind"], r["id"]): r for r in data["rows"]}

        eri = rows[("product", 10)]
        self.assertEqual(eri["qty"], 12)               # 2 retail + 10 wholesale
        self.assertEqual(eri["gift_qty"], 1)           # the gift isn't a sale
        self.assertEqual(eri["revenue"], 620_000)
        self.assertEqual(eri["cost"], 360_000)         # gift costs 0 here (booked in Chiqimlar)
        self.assertEqual(eri["profit"], 260_000)
        self.assertEqual(eri["b2b_qty"], 10)
        self.assertEqual(eri["b2b_revenue"], 500_000)
        self.assertEqual(eri["orders"], 2)
        self.assertTrue(eri["cost_known"])
        # 13 left the shelf in 15 days → 0.87/day; 26 in stock lasts ~30 days.
        self.assertAlmostEqual(eri["per_day"], 0.87, places=2)
        self.assertAlmostEqual(eri["days_left"], 29.9, places=1)

        kokos = rows[("product", 20)]
        self.assertFalse(kokos["cost_known"])          # no cost_price → profit overstated
        self.assertEqual(kokos["days_left"], 0)

        keto_set = rows[("set", 7)]
        self.assertEqual(keto_set["name"], "Keto start to'plami")
        self.assertEqual(keto_set["cost"], 90_000)
        self.assertIsNone(keto_set["per_day"])
        self.assertIsNone(keto_set["days_left"])

        self.assertIn(("product", 99999), rows)        # gone from the catalogue, still sold
        # Biggest revenue first; the unreadable order is skipped, not fatal.
        self.assertEqual(data["rows"][0]["id"], 10)
        self.assertEqual(data["totals"]["revenue"], 620_000 + 45_000 + 150_000 + 30_000 + 10_000)
        self.assertEqual(data["totals"]["orders"], 4)
        self.assertEqual(data["totals"]["days"], 15.0)

    def test_all_time_counts_days_from_the_first_order(self):
        with patch.object(database, "_window_days", wraps=database._window_days) as days:
            data, _ = self._run("all")
        self.assertEqual(days.call_args.args[0], datetime(2026, 10, 12, 8, 0))
        self.assertGreater(data["totals"]["days"], 1)


class PreviousPeriodTest(unittest.TestCase):
    NOW = datetime(2026, 10, 24, 12, 0)          # 17:00 in Tashkent

    def test_rolling_windows(self):
        prev = database.previous_period("7d", self.NOW)
        self.assertEqual(prev, {"start_utc": datetime(2026, 10, 10, 12, 0),
                                "end_utc": datetime(2026, 10, 17, 12, 0)})
        prev = database.previous_period("30d", self.NOW)
        self.assertEqual(prev["end_utc"] - prev["start_utc"], timedelta(days=30))

    def test_today_is_yesterday_up_to_the_same_hour(self):
        prev = database.previous_period("today", self.NOW)
        self.assertEqual(prev, {"start_utc": datetime(2026, 10, 22, 19, 0),
                                "end_utc": datetime(2026, 10, 23, 12, 0)})

    def test_custom_and_all(self):
        prev = database.previous_period({"start": "2026-10-10", "end": "2026-10-24"}, self.NOW)
        start, end = database._period_range({"start": "2026-10-10", "end": "2026-10-24"})
        self.assertEqual(prev["end_utc"], start - timedelta(microseconds=1))
        self.assertEqual(prev["end_utc"] - prev["start_utc"], end - start)
        self.assertIsNone(database.previous_period("all", self.NOW))
        self.assertIsNone(database.previous_period({"start": "x", "end": "y"}, self.NOW))

    def test_exact_window_passes_through_period_range(self):
        w = {"start_utc": datetime(2026, 10, 1), "end_utc": datetime(2026, 10, 2)}
        self.assertEqual(database._period_range(w), (datetime(2026, 10, 1), datetime(2026, 10, 2)))


class _DailyConn:
    def __init__(self):
        self.args = None

    async def fetch(self, sql, *args):
        self.args = args
        return [{"day": date(2026, 10, 20), "orders": 3, "value": 450_000},
                {"day": date(2026, 10, 22), "orders": 1, "value": 99_000}]


class DailySalesTest(unittest.TestCase):
    def test_every_day_is_there_and_short_windows_widen(self):
        conn = _DailyConn()
        with patch.object(database, "pool", _FakePool(conn)):
            days = run(database.get_daily_sales({"start": "2026-10-20", "end": "2026-10-22"}))
        self.assertEqual(len(days), 8)                 # widened to a week
        by_day = {d["date"]: d for d in days}
        self.assertEqual(by_day["2026-10-20"], {"date": "2026-10-20", "orders": 3, "value": 450_000})
        self.assertEqual(by_day["2026-10-21"], {"date": "2026-10-21", "orders": 0, "value": 0})
        self.assertEqual(days[-1]["date"], "2026-10-22")
        self.assertEqual(conn.args[2], database.DISPLAY_TZ_OFFSET_HOURS)

    def test_all_time_is_capped(self):
        conn = _DailyConn()
        with patch.object(database, "pool", _FakePool(conn)):
            days = run(database.get_daily_sales("all", max_days=30))
        self.assertIn(len(days), (31, 32))


class ExcelTest(unittest.TestCase):
    def test_same_rows_and_a_totals_line(self):
        from openpyxl import load_workbook
        conn = _SalesConn()
        with patch.object(database, "pool", _FakePool(conn)), \
                patch.object(database, "get_set_costs", AsyncMock(return_value={7: 90_000})):
            data = run(database.get_product_sales({"start": "2026-10-10", "end": "2026-10-24"}))
        buf, filename = reports.product_sales_excel(data, "10.10.2026 – 24.10.2026")
        self.assertTrue(filename.startswith("ketoshop_mahsulot_savdosi_"))
        ws = load_workbook(BytesIO(buf.getvalue())).active
        self.assertEqual(ws.cell(1, 1).value, "Mahsulot")
        self.assertEqual(ws.cell(2, 1).value, "Eritritol 500 gr")
        self.assertEqual(ws.cell(2, 4).value, 620_000)
        names = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
        self.assertIn("Keto start to'plami (to'plam)", names)
        self.assertIn("Kokos uni ⚠️ tannarx yo'q", names)
        self.assertEqual(names[-1], "JAMI — 10.10.2026 – 24.10.2026")
        self.assertEqual(ws.cell(ws.max_row, 4).value, data["totals"]["revenue"])


class ApiTest(unittest.TestCase):
    def test_period_parsing_and_labels(self):
        import admin_web
        from unittest.mock import MagicMock
        req = MagicMock()
        req.query = {"period": "custom", "start": "2026-10-10", "end": "2026-10-24"}
        period = admin_web._dashboard_period(req)
        self.assertEqual(period, {"start": "2026-10-10", "end": "2026-10-24"})
        self.assertEqual(admin_web._period_label(period), "10.10.2026 – 24.10.2026")
        req.query = {"period": "weird"}
        self.assertEqual(admin_web._dashboard_period(req), "30d")
        self.assertEqual(admin_web._period_label("7d"), "Oxirgi 7 kun")


if __name__ == "__main__":
    unittest.main()
